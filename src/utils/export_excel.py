"""Export movimenti in Excel brandizzato (logo + intestazione)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

EXPORT_COLUMNS = [
    ("data", "Data"),
    ("account", "Conto"),
    ("descrizione", "Descrizione"),
    ("categoria", "Categoria"),
    ("importo", "Importo"),
    ("speciale", "Speciale"),
    ("speciale_mesi", "Mesi ripartizione"),
    ("notes", "Note"),
]

PROJECT_ROOT = Path(__file__).resolve().parents[2]
LOGO_PATH = PROJECT_ROOT / "assets" / "icons" / "ft_logo.png"

DATA_START_ROW = 6


def _prepare_export_frame(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=[label for _, label in EXPORT_COLUMNS])

    frame = df.copy()
    if "data" in frame.columns:
        frame["data"] = pd.to_datetime(
            frame["data"],
            errors="coerce",
            dayfirst=True,
        )
        frame["data"] = frame["data"].dt.strftime("%d/%m/%Y")

    if "speciale" in frame.columns:
        frame["speciale"] = frame["speciale"].fillna(False).map(
            lambda value: "Sì" if bool(value) else "No"
        )
    else:
        frame["speciale"] = "No"

    if "speciale_mesi" in frame.columns:
        frame["speciale_mesi"] = (
            pd.to_numeric(frame["speciale_mesi"], errors="coerce")
            .fillna(0)
            .astype(int)
        )
    else:
        frame["speciale_mesi"] = 0

    columns = [
        source
        for source, _ in EXPORT_COLUMNS
        if source in frame.columns
    ]
    return frame[columns].rename(
        columns={source: label for source, label in EXPORT_COLUMNS}
    )


def movements_to_export_bytes(df: pd.DataFrame) -> bytes:
    export = _prepare_export_frame(df)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Movimenti"

    sheet.merge_cells("B1:D1")
    sheet.merge_cells("B2:D2")
    sheet.merge_cells("B3:D3")

    title_cell = sheet["B1"]
    title_cell.value = "FinanceTracker"
    title_cell.font = Font(name="Calibri", size=18, bold=True, color="0B1220")
    title_cell.alignment = Alignment(vertical="center")

    subtitle_cell = sheet["B2"]
    subtitle_cell.value = "Personal Finance Manager"
    subtitle_cell.font = Font(name="Calibri", size=12, color="64748B")
    subtitle_cell.alignment = Alignment(vertical="center")

    meta_cell = sheet["B3"]
    meta_cell.value = (
        f"Esportato il {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
        f"{len(export)} movimenti"
    )
    meta_cell.font = Font(name="Calibri", size=10, color="94A3B8")

    if LOGO_PATH.exists():
        logo = XLImage(str(LOGO_PATH))
        logo.width = 100
        logo.height = 72
        sheet.add_image(logo, "A1")

    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 18
    sheet.row_dimensions[3].height = 16
    sheet.row_dimensions[4].height = 10

    header_font = Font(name="Calibri", size=11, bold=True, color="0B1220")
    for col_index, column_name in enumerate(export.columns, start=1):
        cell = sheet.cell(row=DATA_START_ROW, column=col_index, value=column_name)
        cell.font = header_font

    for row_index, row in enumerate(export.itertuples(index=False), start=DATA_START_ROW + 1):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)

    for col_index, column_name in enumerate(export.columns, start=1):
        max_length = len(str(column_name))
        for row_index in range(DATA_START_ROW + 1, DATA_START_ROW + 1 + len(export)):
            value = sheet.cell(row=row_index, column=col_index).value
            max_length = max(max_length, len(str(value)) if value is not None else 0)
        sheet.column_dimensions[get_column_letter(col_index)].width = min(
            max(max_length + 2, 12),
            42,
        )

    # Spazio per il logo nella colonna A
    sheet.column_dimensions["A"].width = max(
        sheet.column_dimensions["A"].width or 12,
        14,
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def movements_export_filename(
    *,
    prefix: str = "movimenti",
    account: str | None = None,
    month: str | None = None,
) -> str:
    stamp = datetime.now().strftime("%Y%m%d")
    parts = [prefix]
    if account:
        safe_account = "".join(
            char if char.isalnum() or char in "-_" else "_"
            for char in account.strip()
        ).strip("_") or "conto"
        parts.append(safe_account)
    if month:
        parts.append(month.replace("/", "-"))
    parts.append(stamp)
    return f"{'_'.join(parts)}.xlsx"
